import os
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def download_image(url, filename):
    if os.path.exists(filename):
        return filename
    try:
        r = requests.get(url, timeout=5)
        if r.status_code == 200:
            with open(filename, 'wb') as f:
                f.write(r.content)
            return filename
    except Exception as e:
        print(f"Error downloading {url}: {e}")
    return None

def main():
    prs = Presentation()
    
    # Define layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    blank_slide_layout = prs.slide_layouts[6]
    title_only_layout = prs.slide_layouts[5]

    # Helper to format text
    def add_title(slide, text):
        title = slide.shapes.title
        title.text = text
        for run in title.text_frame.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 102, 204)
            
    def add_bullets(slide, items):
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        for i, item in enumerate(items):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = item
            p.level = 0
            # If item starts with ' -', indent
            if item.startswith(" - "):
                p.text = item.replace(" - ", "")
                p.level = 1

    # Slide 1: Title
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "EASE Air Quality Lab"
    subtitle.text = "Projet de Station Environnementale (Banc de Test)\n\nPrésentation des Architectures Wi-Fi et LoRa"

    # Slide 2: Introduction
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Introduction au Projet EASE")
    add_bullets(slide, [
        "Objectif : Créer une station de mesure de la qualité de l'air intérieure et extérieure.",
        "Phase actuelle : Banc de test multicapteurs pour comparer les technologies.",
        "Modularité : Fonctionne sur Arduino UNO R4 WiFi avec écosystème Grove.",
        "Deux architectures développées :",
        " - Version Réseau Local (Wi-Fi)",
        " - Version Longue Portée (Radio LoRa)"
    ])

    # Slide 3: Alimentation et Interface
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Alimentation et Interface Utilisateur")
    add_bullets(slide, [
        "Autonomie et Flexibilité :",
        " - Le boîtier peut être alimenté par un simple chargeur USB externe.",
        " - Il y a également la place pour intégrer une batterie directement à l'intérieur du boîtier.",
        "Interface Physique :",
        " - Écran OLED intégré pour lire les données en temps réel sans ordinateur.",
        " - Un bouton physique permet de désactiver l'écran pour économiser de l'énergie la nuit ou sur batterie."
    ])

    # Slide 4: Architecture Wi-Fi
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Architecture N°1 : Version Wi-Fi")
    add_bullets(slide, [
        "Structure : Une seule carte Arduino",
        "Fonctionnement : La carte se connecte au réseau Wi-Fi local (box internet ou partage de connexion).",
        "Transmission : Envoi des données via des requêtes HTTP POST au serveur informatique.",
        "Avantage : Simple et direct.",
        "Inconvénient : Portée limitée à la couverture Wi-Fi."
    ])

    # Slide 5: Architecture LoRa
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Architecture N°2 : Version LoRa")
    add_bullets(slide, [
        "Structure : Deux cartes Arduino",
        " - Un 'Émetteur' distant sur batterie, équipé des capteurs et d'une antenne LoRa.",
        " - Un 'Récepteur' branché en USB à l'ordinateur central.",
        "Transmission : Ondes radio 868MHz (portée de plusieurs kilomètres), puis câble USB (Port Série).",
        "Avantage : Peut être placé n'importe où (jardin, forêt, autre bâtiment) sans réseau Wi-Fi."
    ])

    # Slide 6: Aperçu des Boîtiers
    slide = prs.slides.add_slide(title_only_layout)
    add_title(slide, "Aperçu des Boîtiers (Multicapteurs)")
    # Add local images
    try:
        if os.path.exists("boite multicapteurs1.jpg"):
            slide.shapes.add_picture("boite multicapteurs1.jpg", Inches(0.5), Inches(1.5), height=Inches(3.5))
        if os.path.exists("boite multicapteurs2.jpg"):
            slide.shapes.add_picture("boite multicapteurs2.jpg", Inches(5.0), Inches(1.5), height=Inches(3.5))
    except Exception as e:
        print("Erreur image boitier", e)

    # Slide 7: Emplacement Ventilateur et LoRa
    slide = prs.slides.add_slide(title_only_layout)
    add_title(slide, "Aperçu : Ventilation et Boîtier LoRa")
    try:
        if os.path.exists("emplacement ventilateur.jpg"):
            slide.shapes.add_picture("emplacement ventilateur.jpg", Inches(0.5), Inches(1.5), height=Inches(3.5))
        if os.path.exists("boite LoRa.jpg"):
            slide.shapes.add_picture("boite LoRa.jpg", Inches(5.0), Inches(1.5), height=Inches(3.5))
    except Exception as e:
        print("Erreur image ventilation/lora", e)

    # Slide 8: Récupération des Données
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Récupération et Traitement des Données")
    add_bullets(slide, [
        "Un serveur informatique codé en Python collecte toutes les données.",
        "Technologie : Framework Flask et serveur de production Waitress.",
        "Mécanisme de récupération :",
        " - Wi-Fi : Le serveur écoute sur le port 5000 les paquets HTTP envoyés par l'Arduino.",
        " - LoRa : Le serveur écoute le port COM (USB) du Récepteur via la bibliothèque PySerial.",
        "Visualisation : Un tableau de bord complet (Dashboard Web) affiche les graphiques en temps réel et l'historique."
    ])

    # Slide 9: Format des Données (CSV)
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Format de Sauvegarde")
    add_bullets(slide, [
        "Les données sont stockées sous forme de fichier CSV (tableur).",
        "Structure : 23 colonnes mesurées toutes les 10 secondes.",
        "Données collectées :",
        " - Horodatage (Date et Heure)",
        " - Paramètres ambiants (Température, Humidité, Pression)",
        " - Qualité de l'air (CO2, COV, eCO2, HCHO)",
        " - Particules fines (PM1.0, PM2.5, PM10)",
        "Horodatage Autonome :",
        " - Utilisation d'une horloge RTC (PCF85063) sur la carte pour une datation fiable même hors connexion internet."
    ])

    # Sensor Images downloading
    sensors = {
        "dht20.jpg": "https://files.seeedstudio.com/wiki/Grove-Temperature-Humidity-Sensor-V2.0-DHT20/img/main.jpg",
        "bme680.jpg": "https://files.seeedstudio.com/wiki/Grove-Temperature_Humidity_Pressure_Gas_Sensor_BME680/img/main.jpg",
        "scd30.jpg": "https://files.seeedstudio.com/wiki/Grove-CO2_Temperature_Humidity_Sensor-SCD30/img/main.jpg",
        "hm3301.jpg": "https://files.seeedstudio.com/wiki/Grove-Laser_PM2.5_Sensor-HM3301/img/main.jpg",
        "multi.jpg": "https://files.seeedstudio.com/wiki/Grove-Multichannel-Gas-Sensor-V2/img/main.jpg"
    }
    
    for fname, url in sensors.items():
        download_image(url, fname)

    # Slide 10: Capteurs Température & Environnement
    slide = prs.slides.add_slide(title_only_layout)
    add_title(slide, "Capteurs : Environnement Général")
    # Text box
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(2))
    tf = txBox.text_frame
    tf.text = "Deux capteurs pour comparer l'ambiance thermique et la pression :"
    p = tf.add_paragraph()
    p.text = "- DHT20 : Référence neutre pour la Température et l'Humidité."
    p = tf.add_paragraph()
    p.text = "- BME680 : Température, Humidité, Pression Atmo, et résistance des Gaz (COV)."
    # Images
    if os.path.exists("dht20.jpg"):
        slide.shapes.add_picture("dht20.jpg", Inches(1.0), Inches(3.5), height=Inches(2.5))
    if os.path.exists("bme680.jpg"):
        slide.shapes.add_picture("bme680.jpg", Inches(5.5), Inches(3.5), height=Inches(2.5))

    # Slide 11: Capteurs CO2 et Particules
    slide = prs.slides.add_slide(title_only_layout)
    add_title(slide, "Capteurs : CO2 & Particules")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(2))
    tf = txBox.text_frame
    tf.text = "Les capteurs de spécialité :"
    p = tf.add_paragraph()
    p.text = "- SCD30 (NDIR) : Mesure optique infrarouge du vrai CO2 avec étalonnage automatique."
    p = tf.add_paragraph()
    p.text = "- HM3301 (Laser) : Compteur optique de micro-particules (PM1.0, PM2.5, PM10)."
    
    if os.path.exists("scd30.jpg"):
        slide.shapes.add_picture("scd30.jpg", Inches(1.0), Inches(3.5), height=Inches(2.5))
    if os.path.exists("hm3301.jpg"):
        slide.shapes.add_picture("hm3301.jpg", Inches(5.5), Inches(3.5), height=Inches(2.5))

    # Slide 12: Capteurs Gaz Chimiques
    slide = prs.slides.add_slide(title_only_layout)
    add_title(slide, "Capteurs : Gaz Chimiques (MOX)")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(2))
    tf = txBox.text_frame
    tf.text = "Utilisation d'oxydes métalliques (nécessite un préchauffage) :"
    p = tf.add_paragraph()
    p.text = "- Multichannel V2 : Tendances pour NO2, CO, Éthanol, VOC."
    p = tf.add_paragraph()
    p.text = "- WSP2110 : Détection spécifique du Formaldéhyde (HCHO)."
    
    if os.path.exists("multi.jpg"):
        slide.shapes.add_picture("multi.jpg", Inches(3.5), Inches(3.5), height=Inches(2.5))

    # Slide 13: Conclusions et Évolutions
    slide = prs.slides.add_slide(bullet_slide_layout)
    add_title(slide, "Évolutions Futures")
    add_bullets(slide, [
        "Intégration d'une ventilation active :",
        " - Problème actuel : Effet de serre et auto-échauffement des capteurs (Infrarouge et MOX).",
        " - Solution : Ajout d'un ventilateur (ex: Noctua 40mm) sur l'emplacement prévu pour assurer un flux d'air naturel.",
        "Création de la version finale :",
        " - Sélection du meilleur capteur de chaque catégorie suite aux analyses de données.",
        " - Réduction de la taille et de la consommation énergétique."
    ])

    # Ajouter le logo sur toutes les slides
    logo_path = r"c:\Users\sylva\Documents\! projet climat\station_complete_lora\static\ease_logo.png"
    if os.path.exists(logo_path):
        for s in prs.slides:
            s.shapes.add_picture(logo_path, Inches(8.5), Inches(0.2), height=Inches(0.8))

    prs.save("Presentation_Projet_EASE.pptx")
    print("Fichier Presentation_Projet_EASE.pptx généré avec succès !")

if __name__ == "__main__":
    main()
